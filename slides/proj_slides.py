"""Course-project section. Starts with setup and the handover env config."""

from __future__ import annotations

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

from pathlib import Path

from build_rl_section import INK, MUTED, NAVY, add_muli_run
from sim_slides import _box, _contain_picture, _emu, _fresh, _title

NOTE = "Notebook: hri_handover_tutorial.ipynb"
ROOT = Path(__file__).resolve().parent
PROJ = ROOT / "images" / "proj"


def _section(prs, exemplar: int, title: str, part: str) -> None:
    master = prs.slides[exemplar].slide_layout.slide_master
    layout = None
    for candidate in master.slide_layouts:
        c_sld = candidate._element.find(qn("p:cSld"))
        if c_sld is not None and c_sld.get("name") == "Section Header":
            layout = candidate
            break
    if layout is None:
        raise SystemExit("No light Section Header layout.")
    slide = prs.slides.add_slide(layout)
    for shape in slide.placeholders:
        kind = shape.placeholder_format.type
        if kind == PP_PLACEHOLDER.TITLE:
            shape.text_frame.clear()
            add_muli_run(shape.text_frame.paragraphs[0], title, "Muli ExtraBold", 32, True)
        elif kind == PP_PLACEHOLDER.BODY:
            shape.text_frame.clear()
            add_muli_run(shape.text_frame.paragraphs[0], part, "Muli", 18, False)


def _note(slide) -> None:
    _box(slide, 1.35, 5.12, 5.85, 0.28, [NOTE], size=9, color=MUTED, align=PP_ALIGN.RIGHT)


def _panel(slide, x, y, w, h) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, _emu(x), _emu(y), _emu(w), _emu(h))
    shape.adjustments[0] = 0.06
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0xF4, 0xF7, 0xFA)
    shape.line.color.rgb = RGBColor(0xD5, 0xDC, 0xE3)
    shape.line.width = Pt(1)


def _code(slide, x, y, w, h, lines, *, size=13) -> None:
    shape = slide.shapes.add_textbox(_emu(x), _emu(y), _emu(w), _emu(h))
    tf = shape.text_frame
    tf.word_wrap = False
    tf.auto_size = None
    tf._txBody.find(qn("a:bodyPr")).set("anchor", "t")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(3)
        add_muli_run(p, line, "Consolas", size, False, INK)
    return shape


def _setup(prs, exemplar: int) -> None:
    slide = _fresh(prs, exemplar)
    _title(slide, "Set up")
    _box(slide, 0.42, 0.82, 4.50, 0.36, [("Hardware", True, 18, NAVY)])
    _box(
        slide,
        0.42,
        1.14,
        4.50,
        1.45,
        [
            [("•  macOS, Linux, or Windows", False, 16, INK)],
            [("•  8 GB RAM, a few GB of disk", False, 16, INK)],
            [("•  GPU optional; CPU is enough", False, 16, INK)],
            [("      to open the notebook", False, 15, INK)],
        ],
    )
    _box(slide, 5.10, 0.82, 4.50, 0.36, [("Get the repo", True, 18, NAVY)])
    _box(
        slide,
        5.10,
        1.14,
        4.50,
        1.45,
        [
            [("•  Clone CourseProj from GitHub", False, 16, INK)],
            [("•  Conda, Python 3.10", False, 16, INK)],
            [("•  pip installs local robosuite", False, 16, INK)],
            [("•  Then open the notebook", False, 16, INK)],
        ],
    )
    _panel(slide, 0.42, 2.62, 9.18, 2.00)
    _code(
        slide,
        0.62,
        2.74,
        8.80,
        1.78,
        [
            "git clone https://github.com/RentonRen/CourseProj.git",
            "cd CourseProj",
            "conda create -n courseproj python=3.10",
            "conda activate courseproj",
            "pip install -r requirements.txt",
            "jupyter notebook hri_handover_tutorial.ipynb",
        ],
        size=14,
    )
    _box(
        slide,
        0.42,
        4.88,
        8.8,
        0.24,
        ["Windows: create C:\\tmp, copy mujoco.dll into robosuite/utils, set MUJOCO_GL=wgl."],
        size=11,
        color=MUTED,
    )
    _note(slide)


def _robosuite(prs, exemplar: int) -> None:
    slide = _fresh(prs, exemplar)
    _title(slide, "Part 0")
    _box(
        slide,
        0.42,
        0.82,
        4.75,
        2.55,
        [
            [("A robotics library built on MuJoCo.", False, 15, INK)],
            [("It provides:", False, 15, INK)],
            [("•  Ready-made robot arms", False, 15, INK)],
            [("      Panda, XArm7, UR5e, …", False, 14, INK)],
            [("•  Ready-made manipulation tasks", False, 15, INK)],
            [("      Lift, Pick-and-Place, Assembly, …", False, 14, INK)],
            [("•  Cameras, sensors, and a standard", False, 15, INK)],
            [("      observation / action interface", False, 14, INK)],
            [("•  Built-in scenes to test as-is,", False, 15, INK)],
            [("      or change a little and retest", False, 14, INK)],
        ],
    )
    _panel(slide, 0.42, 3.40, 4.75, 1.28)
    _code(
        slide,
        0.55,
        3.52,
        4.50,
        1.10,
        [
            "obs = env.reset()",
            "obs, reward, done, info = env.step(action)",
        ],
        size=12,
    )
    _contain_picture(slide, PROJ / "part0_imports.png", 5.25, 0.86, 4.35, 3.55)
    _box(
        slide,
        5.25,
        4.48,
        4.35,
        0.24,
        ["Part 0 of hri_handover_tutorial.ipynb"],
        size=10,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
    )
    _box(
        slide,
        0.42,
        4.78,
        6.2,
        0.22,
        ["Zhu et al., robosuite. arXiv:2009.12293, 2020. 2017 Stanford SVL; now ARISE."],
        size=10,
        color=MUTED,
    )
    _note(slide)


def _env_config(prs, exemplar: int) -> None:
    slide = _fresh(prs, exemplar)
    _title(slide, "Part 1")
    _panel(slide, 0.40, 0.88, 5.35, 3.95)
    _code(
        slide,
        0.55,
        1.00,
        5.10,
        3.75,
        [
            "env = suite.make(",
            '    env_name="HumanRobotHandover",',
            '    robots="XArm7Handover",',
            "    has_renderer=False,",
            "    has_offscreen_renderer=True,",
            "    use_camera_obs=True,",
            "    horizon=300,",
            "    control_freq=20,",
            "    reward_shaping=False,",
            "    initialization_noise=None,",
            "    seed=0,",
            "    hard_reset=False,",
            ")",
            "env.start_holding = True",
        ],
        size=13,
    )
    _box(
        slide,
        5.90,
        0.92,
        3.70,
        3.95,
        [
            [("•  One XArm7, hammer in hand", False, 16, INK)],
            [("•  20 Hz, 300 steps, 15 s", False, 16, INK)],
            [("•  Headless images in memory", False, 16, INK)],
            [("•  Cameras on for the tour", False, 16, INK)],
            [("•  Sparse reward at first", False, 16, INK)],
            [("•  seed=0, held start pose", False, 16, INK)],
            [("•  Later: cameras off,", False, 16, INK)],
            [("      reward_shaping=True", False, 16, INK)],
        ],
    )
    _note(slide)


def add_project_slides(prs, exemplar: int) -> None:
    _section(prs, exemplar, "Course project", "Part 3 of 3")
    _setup(prs, exemplar)
    _robosuite(prs, exemplar)
    _env_config(prs, exemplar)
